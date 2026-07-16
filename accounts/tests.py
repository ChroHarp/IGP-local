from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.core.files.storage import InMemoryStorage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from django.test import TestCase
from django.urls import reverse

from .importers import import_basic_students
from .models import AwardRecord, FamilyMember, Guardian, InitialIGPProfile, ProgramDocument, Student, StudentStaffAssignment, Teacher
from .policies import approved_google_user_for_email, can_manage_learning_outcomes, visible_students_for


class ProjectSmokeTests(TestCase):
    def test_health_check(self):
        response = self.client.get(reverse("health"), secure=True)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, b"ok")

    def test_custom_user_model_is_active(self):
        user = get_user_model().objects.create_user(
            username="phase-one-user",
            email="phase-one@example.edu.tw",
            password="safe-test-password",
        )

        self.assertEqual(user.email, "phase-one@example.edu.tw")


class AccountLoginTests(TestCase):
    def test_approved_local_user_can_open_login_and_log_in(self):
        user = get_user_model().objects.create_user(
            username="joyce", email="joyce@example.edu.tw", password="safe-test-password",
            is_approved=True, is_active=True,
        )

        login_page = self.client.get(reverse("account_login"))
        self.assertEqual(login_page.status_code, 200)
        self.assertContains(login_page, "auth-shell")
        self.assertContains(login_page, "IGP 本地管理")
        response = self.client.post(reverse("account_login"), {"login": user.username, "password": "safe-test-password"})

        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.client.post(reverse("account_logout")).status_code, 302)


class PhaseOneModelTests(TestCase):
    def setUp(self):
        self.student = Student.objects.create(full_name="測試學生")
        self.teacher = get_user_model().objects.create_user(
            username="teacher",
            email="teacher@example.edu.tw",
            password="safe-test-password",
            is_approved=True,
        )
        self.teacher_record = Teacher.objects.create(
            full_name="測試教師", account=self.teacher
        )

    def test_only_one_primary_guardian_is_allowed(self):
        Guardian.objects.create(
            student=self.student,
            relationship=Guardian.Relationship.LEGAL_GUARDIAN,
            full_name="測試家長一",
            is_primary=True,
        )

        with self.assertRaises(IntegrityError), transaction.atomic():
            Guardian.objects.create(
                student=self.student,
                relationship=Guardian.Relationship.OTHER,
                full_name="測試家長二",
                is_primary=True,
            )

    def test_assignment_rejects_an_end_date_before_its_start_date(self):
        assignment = StudentStaffAssignment(
            student=self.student,
            staff=self.teacher_record,
            role=StudentStaffAssignment.Role.COURSE_TEACHER,
            start_date=date.today(),
            end_date=date.today() - timedelta(days=1),
        )

        with self.assertRaises(ValidationError):
            assignment.full_clean()

    def test_only_one_current_case_manager_is_allowed(self):
        other_account = get_user_model().objects.create_user(
            username="other-case-manager",
            email="other-case-manager@example.edu.tw",
            password="safe-test-password",
        )
        other_teacher = Teacher.objects.create(full_name="other", account=other_account)
        StudentStaffAssignment.objects.create(
            student=self.student,
            staff=self.teacher_record,
            role=StudentStaffAssignment.Role.CASE_MANAGER,
        )
        duplicate = StudentStaffAssignment(
            student=self.student,
            staff=other_teacher,
            role=StudentStaffAssignment.Role.CASE_MANAGER,
        )

        with self.assertRaises(ValidationError):
            duplicate.full_clean()

    def test_google_policy_requires_an_approved_active_local_user(self):
        self.assertEqual(
            approved_google_user_for_email("TEACHER@example.edu.tw"),
            self.teacher,
        )

        self.teacher.is_approved = False
        self.teacher.save(update_fields=["is_approved"])
        self.assertIsNone(approved_google_user_for_email("teacher@example.edu.tw"))


    def test_phase_two_plan_links_learning_outcomes_to_a_student(self):
        from .models import CoursePlan, IGPPlan, LearningOutcome, SemesterPlan

        plan = IGPPlan.objects.create(student=self.student, academic_year="114", overall_goal="research")
        semester = SemesterPlan.objects.create(igp_plan=plan, semester=1, goals="project")
        course = CoursePlan.objects.create(semester_plan=semester, course_name="science", goals="question")
        outcome = LearningOutcome.objects.create(course_plan=course, outcome="record")

        self.assertEqual(outcome.course_plan.semester_plan.igp_plan.student, self.student)


class StudentAccessTests(TestCase):
    def setUp(self):
        self.visible_student = Student.objects.create(full_name="可見學生")
        self.hidden_student = Student.objects.create(full_name="不可見學生")
        self.case_manager = self.create_user("case-manager", get_user_model().Role.CASE_MANAGER)
        self.course_teacher = self.create_user("course-teacher", get_user_model().Role.COURSE_TEACHER)
        self.system_admin = self.create_user("system-admin", get_user_model().Role.SYSTEM_ADMIN)
        self.school_lead = self.create_user("school-lead", get_user_model().Role.SPECIAL_EDUCATION_LEAD)
        self.case_manager_record = Teacher.objects.create(
            full_name="個管教師", account=self.case_manager
        )
        StudentStaffAssignment.objects.create(
            student=self.visible_student,
            staff=self.case_manager_record,
            role=StudentStaffAssignment.Role.CASE_MANAGER,
        )

    def create_user(self, username, role):
        return get_user_model().objects.create_user(
            username=username,
            email=f"{username}@example.edu.tw",
            password="safe-test-password",
            role=role,
            is_approved=True,
            is_staff=True,
        )

    def test_case_manager_only_sees_assigned_students(self):
        self.assertQuerySetEqual(visible_students_for(self.case_manager), [self.visible_student])

    def test_course_teacher_and_system_admin_see_no_student_data_by_default(self):
        self.assertQuerySetEqual(visible_students_for(self.course_teacher), [])
        self.assertQuerySetEqual(visible_students_for(self.system_admin), [])

    def test_special_education_lead_sees_all_students(self):
        self.assertQuerySetEqual(
            visible_students_for(self.school_lead),
            [self.visible_student, self.hidden_student],
            ordered=False,
        )

    def test_admin_list_never_exposes_an_unassigned_student(self):
        self.client.force_login(self.case_manager)

        response = self.client.get(reverse("admin:accounts_student_changelist"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "可見學生")
        self.assertNotContains(response, "不可見學生")
    def test_student_change_form_groups_data_into_tabs(self):
        InitialIGPProfile.objects.create(
            student=self.visible_student,
            family_culture="重視閱讀",
            cognitive_strengths="理解力佳",
        )
        AwardRecord.objects.create(
            student=self.visible_student,
            activity_name="科展",
            award="第一名",
        )
        self.client.force_login(self.school_lead)

        response = self.client.get(reverse("admin:accounts_student_change", args=[self.visible_student.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "基本資料")
        self.assertContains(response, "家庭與聯絡")
        self.assertContains(response, "學生能力分析")
        self.assertContains(response, "得獎與紀錄")
        self.assertContains(response, "重視閱讀")
        self.assertContains(response, "科展")
    def test_igp_plans_default_to_the_school_academic_year(self):
        from .models import IGPPlan, SchoolSetting

        SchoolSetting.objects.create(name="Test School", academic_year="115")
        IGPPlan.objects.create(student=self.visible_student, academic_year="114", overall_goal="obsolete-goal-marker")
        IGPPlan.objects.create(student=self.hidden_student, academic_year="115", overall_goal="current")
        self.client.force_login(self.school_lead)

        response = self.client.get(reverse("admin:accounts_igpplan_changelist"))

        self.assertContains(response, "current")
        self.assertNotContains(response, "obsolete-goal-marker")

    def test_selected_semester_plan_can_open_the_copy_action(self):
        from .models import IGPPlan, SemesterPlan

        plan = IGPPlan.objects.create(student=self.visible_student, academic_year="115", overall_goal="goal")
        semester = SemesterPlan.objects.create(igp_plan=plan, semester=1, goals="semester goal")
        self.client.force_login(self.school_lead)

        response = self.client.post(
            reverse("admin:accounts_semesterplan_changelist"),
            {"action": "copy_selected_plan", "_selected_action": semester.pk},
        )

        self.assertRedirects(response, reverse("admin:accounts_semesterplan_copy", args=[semester.pk]))

    def test_special_education_lead_can_open_the_import_page(self):
        self.client.force_login(self.school_lead)

        response = self.client.get(reverse("admin:accounts_student_import"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "匯入學生基礎資料")
    def test_superuser_can_open_an_empty_student_list(self):
        StudentStaffAssignment.objects.all().delete()
        Student.objects.all().delete()
        get_user_model().objects.all().delete()
        Student.objects.all().delete()
        superuser = get_user_model().objects.create_superuser(
            username="break-glass",
            email="break-glass@example.edu.tw",
            password="safe-test-password",
        )
        self.client.force_login(superuser)

        response = self.client.get(reverse("admin:accounts_student_changelist"))

        self.assertEqual(response.status_code, 200)

class BatchImportTests(TestCase):
    def workbook(self, rows):
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["學生姓名", "性別", "年級", "班別", "座號", "出生年月日", "法定代理人", "法定代理人連絡電話 (手機)"])
        for row in rows:
            worksheet.append(row)
        data = BytesIO()
        workbook.save(data)
        data.seek(0)
        return data

    def test_preview_does_not_write_student_data(self):
        result = import_basic_students(
            self.workbook([["預覽學生", "女", 7, "701", 3, date(2013, 1, 2), "預覽家長", "0912345678"]]),
            apply=False,
        )

        self.assertTrue(result.is_valid)
        self.assertEqual(result.create_count, 1)
        self.assertFalse(Student.objects.filter(full_name="預覽學生").exists())

    def test_apply_creates_student_and_primary_guardian(self):
        result = import_basic_students(
            self.workbook([["匯入學生", "男", 7, "701", 4, date(2013, 2, 3), "匯入家長", "0987654321"]]),
            apply=True,
        )

        student = Student.objects.get(full_name="匯入學生")
        self.assertTrue(result.is_valid)
        self.assertEqual(student.guardians.get().full_name, "匯入家長")
        self.assertTrue(student.guardians.get().is_primary)
        self.assertTrue(InitialIGPProfile.objects.filter(student=student).exists())

    def test_import_maps_igp_profile_family_award_and_multiple_needs(self):
        from openpyxl import Workbook

        headers = [
            "學生姓名", "性別", "年級", "班別", "是否有雙重特教需求", "父親姓名", "父親畢業科系",
            "父親專長", "父親連絡電話", "父親服務機關", "家庭文化特質", "主要照顧者", "主要協助學習者",
            "家庭經濟狀況", "照顧者管教態度", "與家人互動情形", "第一階段-數學性向測驗分數",
            "科學興趣", "認知優勢特質", "情意弱勢特質", "獲獎日期", "競賽/活動名稱", "主辦單位", "獎項", "得獎類型", "填寫者",
        ]
        row = [
            "完整匯入學生", "女", 7, "701", "是", "王父", "工程", "程式設計", "0911", "科技公司",
            "一般生", "父親, 母親", "母親", "小康", "民主式", "5", "97", "數學, 天文",
            "記憶能力, 理解能力", "情緒控制", "2025-01-21", "科展", "教育局", "第一名", "團體獎", "家長觀察",
        ]
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(headers)
        worksheet.append(row)
        data = BytesIO()
        workbook.save(data)
        data.seek(0)

        result = import_basic_students(data, apply=True)

        student = Student.objects.get(full_name="完整匯入學生")
        profile = student.initial_igp_profile
        self.assertTrue(result.is_valid)
        self.assertTrue(student.has_multiple_special_education_needs)
        self.assertEqual(profile.family_culture, "一般生")
        self.assertEqual(profile.primary_caregiver, "父親, 母親")
        self.assertEqual(profile.cognitive_strengths, "記憶能力, 理解能力")
        self.assertEqual(FamilyMember.objects.get(student=student, relationship="父親").specialty, "程式設計")
        self.assertEqual(AwardRecord.objects.get(student=student).activity_name, "科展")

    def test_backfill_populates_existing_raw_response(self):
        student = Student.objects.create(full_name="回填學生")
        InitialIGPProfile.objects.create(
            student=student,
            raw_response={"是否有雙重特教需求": "是", "科學興趣": "數學, 生物", "填寫者": "學生自評"},
        )

        call_command("backfill_igp_profiles")

        student.refresh_from_db()
        student.initial_igp_profile.refresh_from_db()
        self.assertTrue(student.has_multiple_special_education_needs)
        self.assertEqual(student.initial_igp_profile.science_interests, "數學, 生物")
        self.assertEqual(student.initial_igp_profile.completed_by, "學生自評")
    def test_duplicate_rows_are_rejected_without_writing(self):
        result = import_basic_students(
            self.workbook([
                ["重複學生", "女", 7, "701", 1, date(2013, 3, 4), "家長", ""],
                ["重複學生", "女", 7, "701", 2, date(2013, 3, 4), "家長", ""],
            ]),
            apply=True,
        )

        self.assertFalse(result.is_valid)
        self.assertFalse(Student.objects.filter(full_name="重複學生").exists())


class ProgramDocumentTests(TestCase):
    def setUp(self):
        self.document_field = ProgramDocument._meta.get_field("document_file")
        self.original_storage = self.document_field.storage
        self.document_field.storage = InMemoryStorage()
        self.addCleanup(setattr, self.document_field, "storage", self.original_storage)

    def create_document(self, filename, content):
        return ProgramDocument(
            document_type=ProgramDocument.DocumentType.COURSE_PLAN,
            title="測試文件",
            document_file=SimpleUploadedFile(filename, content),
        )

    def test_valid_pdf_is_accepted_and_fake_pdf_is_rejected(self):
        self.create_document("plan.pdf", b"%PDF-1.7\ncontent").full_clean()
        with self.assertRaises(ValidationError):
            self.create_document("fake.pdf", b"not a pdf").full_clean()

    def test_valid_docx_requires_word_document_xml(self):
        data = BytesIO()
        with ZipFile(data, "w") as archive:
            archive.writestr("word/document.xml", "<w:document />")
        self.create_document("plan.docx", data.getvalue()).full_clean()

    def test_download_is_denied_to_system_admin_and_allowed_to_course_teacher(self):
        document = ProgramDocument.objects.create(
            document_type=ProgramDocument.DocumentType.TIMETABLE,
            title="課表",
            document_file=SimpleUploadedFile("timetable.pdf", b"%PDF-1.7\ncontent"),
        )
        system_admin = get_user_model().objects.create_user(
            username="document-system-admin",
            email="document-system-admin@example.edu.tw",
            password="safe-test-password",
            role=get_user_model().Role.SYSTEM_ADMIN,
            is_staff=True,
        )
        course_teacher = get_user_model().objects.create_user(
            username="document-course-teacher",
            email="document-course-teacher@example.edu.tw",
            password="safe-test-password",
            role=get_user_model().Role.COURSE_TEACHER,
            is_staff=True,
        )

        self.client.force_login(system_admin)
        denied = self.client.get(reverse("program-document-download", args=[document.public_id]))
        self.assertEqual(denied.status_code, 403)

        self.client.force_login(course_teacher)
        allowed = self.client.get(reverse("program-document-download", args=[document.public_id]))
        self.assertEqual(allowed.status_code, 200)









class TeacherAssignmentBoardTests(TestCase):
    def setUp(self):
        self.lead = get_user_model().objects.create_user(
            username="lead", email="lead@example.edu.tw", password="safe-test-password",
            role=get_user_model().Role.SPECIAL_EDUCATION_LEAD, is_approved=True, is_staff=True,
        )
        self.teacher_account = get_user_model().objects.create_user(
            username="course", email="course@example.edu.tw", password="safe-test-password",
            role=get_user_model().Role.COURSE_TEACHER, is_approved=True, is_staff=True,
        )
        self.teacher = Teacher.objects.create(
            full_name="course", account=self.teacher_account,
        )
        self.first_student = Student.objects.create(full_name="第一位學生")
        self.second_student = Student.objects.create(full_name="第二位學生")
        self.url = reverse("admin:accounts_studentstaffassignment_teacher", args=[self.teacher.pk])

    def test_lead_can_assign_case_manager_students_and_preserve_unchecked_history(self):
        self.client.force_login(self.lead)
        board = self.client.get(reverse("admin:accounts_studentstaffassignment_changelist"))
        self.assertEqual(board.status_code, 200)
        self.assertContains(board, "course")

        self.client.post(self.url, {"account": self.teacher_account.pk, "case_manager_students": [self.first_student.pk, self.second_student.pk]})
        self.assertEqual(StudentStaffAssignment.objects.filter(staff=self.teacher, is_active=True).count(), 2)

        self.client.post(self.url, {"account": self.teacher_account.pk, "case_manager_students": [self.first_student.pk]})
        self.assertTrue(StudentStaffAssignment.objects.get(staff=self.teacher, student=self.first_student).is_active)
        self.assertFalse(StudentStaffAssignment.objects.get(staff=self.teacher, student=self.second_student).is_active)

    def test_homeroom_and_case_manager_roles_are_saved_independently(self):
        self.client.force_login(self.lead)

        self.client.post(self.url, {
            "account": self.teacher_account.pk,
            "homeroom_students": [self.first_student.pk, self.second_student.pk],
            "case_manager_students": [self.second_student.pk],
        })

        active = StudentStaffAssignment.objects.filter(staff=self.teacher, is_active=True)
        self.assertEqual(active.count(), 3)
        self.assertTrue(active.filter(student=self.second_student, role=StudentStaffAssignment.Role.HOMEROOM_TEACHER).exists())
        self.assertTrue(active.filter(student=self.second_student, role=StudentStaffAssignment.Role.CASE_MANAGER).exists())
        self.assertFalse(active.filter(role=StudentStaffAssignment.Role.COURSE_TEACHER).exists())
        self.client.post(self.url, {
            "account": self.teacher_account.pk,
            "homeroom_students": [self.first_student.pk],
            "case_manager_students": [self.second_student.pk],
        })

        self.assertFalse(StudentStaffAssignment.objects.get(student=self.second_student, role=StudentStaffAssignment.Role.HOMEROOM_TEACHER).is_active)
        self.assertTrue(active.get(student=self.second_student, role=StudentStaffAssignment.Role.CASE_MANAGER).is_active)

    def test_case_manager_assignment_shows_validation_error_when_student_already_has_one(self):
        other_account = get_user_model().objects.create_user(
            username="other-case-manager", email="other-case-manager@example.edu.tw",
            password="safe-test-password", role=get_user_model().Role.CASE_MANAGER,
            is_approved=True, is_staff=True,
        )
        other_teacher = Teacher.objects.create(full_name="既有個管教師", account=other_account)
        StudentStaffAssignment.objects.create(
            student=self.first_student,
            staff=other_teacher,
            role=StudentStaffAssignment.Role.CASE_MANAGER,
        )
        self.client.force_login(self.lead)

        response = self.client.post(self.url, {
            "account": self.teacher_account.pk,
            "case_manager_students": [self.first_student.pk],
        })

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "第一位學生 已是 既有個管教師 的個案")
        self.assertFalse(StudentStaffAssignment.objects.filter(
            student=self.first_student,
            staff=self.teacher,
            role=StudentStaffAssignment.Role.CASE_MANAGER,
            is_active=True,
        ).exists())
    def test_course_students_cannot_be_assigned_from_teacher_board(self):
        self.client.force_login(self.lead)

        self.client.post(self.url, {
            "account": self.teacher_account.pk,
            "course_students": [self.first_student.pk],
        })

        self.assertFalse(StudentStaffAssignment.objects.filter(
            staff=self.teacher,
            role=StudentStaffAssignment.Role.COURSE_TEACHER,
        ).exists())
    def test_lead_can_create_a_teacher_without_an_account(self):
        self.client.force_login(self.lead)
        response = self.client.post(
            reverse("admin:accounts_studentstaffassignment_teacher_add"),
            {"full_name": "新任教師"},
        )

        teacher = Teacher.objects.get(full_name="新任教師")
        self.assertRedirects(response, reverse("admin:accounts_studentstaffassignment_teacher", args=[teacher.pk]))
        self.assertIsNone(teacher.account)

    def test_account_can_be_linked_later(self):
        account = get_user_model().objects.create_user(
            username="later", email="later@example.edu.tw", password="safe-test-password",
            role=get_user_model().Role.COURSE_TEACHER, is_approved=True, is_staff=True,
        )
        teacher = Teacher.objects.create(full_name="稍後綁定")
        self.client.force_login(self.lead)

        response = self.client.post(
            reverse("admin:accounts_studentstaffassignment_teacher", args=[teacher.pk]),
            {"account": account.pk},
        )

        teacher.refresh_from_db()
        self.assertRedirects(response, reverse("admin:accounts_studentstaffassignment_changelist"))
        self.assertEqual(teacher.account, account)

    def test_board_edits_case_and_homeroom_assignments_and_shows_course_students_read_only(self):
        from .models import CoursePlan, IGPPlan, SemesterPlan

        annual = IGPPlan.objects.create(student=self.first_student, academic_year="115", overall_goal="goal")
        semester = SemesterPlan.objects.create(igp_plan=annual, semester=1, goals="semester goal")
        CoursePlan.objects.create(
            semester_plan=semester,
            course_name="Mathematics",
            teacher=self.teacher,
            goals="course goal",
        )
        self.client.force_login(self.lead)

        board = self.client.get(reverse("admin:accounts_studentstaffassignment_changelist"))
        assignment = self.client.get(self.url)

        self.assertContains(board, "新增教師")
        self.assertContains(board, "指派學生")
        self.assertContains(board, "任課教師")
        self.assertContains(assignment, 'type="checkbox"', count=4)
        self.assertContains(assignment, "任課學生（由課程計畫自動帶入）")
        self.assertContains(assignment, self.first_student.full_name)
        self.assertNotContains(assignment, 'name="course_students"')

    def test_lead_can_change_a_users_role_and_approval(self):
        account = get_user_model().objects.create_user(
            username="editable-user", email="editable@example.edu.tw", password="safe-test-password",
            role=get_user_model().Role.VIEWER, is_staff=True,
        )
        self.client.force_login(self.lead)

        response = self.client.post(
            reverse("admin:accounts_user_change", args=[account.pk]),
            {
                "username": account.username,
                "email": account.email,
                "role": get_user_model().Role.COURSE_TEACHER,
                "is_approved": "on",
                "is_active": "on",
                "is_staff": "on",
                "_save": "Save",
            },
        )

        account.refresh_from_db()
        self.assertRedirects(response, reverse("admin:accounts_user_changelist"))
        self.assertEqual(account.role, get_user_model().Role.COURSE_TEACHER)
        self.assertTrue(account.is_approved)

    def test_teacher_can_be_created_with_an_approved_account(self):
        account = get_user_model().objects.create_user(
            username="matchable", email="matchable@example.edu.tw", password="safe-test-password",
            role=get_user_model().Role.COURSE_TEACHER, is_approved=True, is_staff=True,
        )
        self.client.force_login(self.lead)

        response = self.client.post(
            reverse("admin:accounts_studentstaffassignment_teacher_add"),
            {"full_name": "Matched Teacher", "account": account.pk},
        )

        teacher = Teacher.objects.get(full_name="Matched Teacher")
        self.assertRedirects(response, reverse("admin:accounts_studentstaffassignment_teacher", args=[teacher.pk]))
        self.assertEqual(teacher.account, account)

    def test_user_admin_requires_email_instead_of_saving_an_empty_unique_value(self):
        self.client.force_login(self.lead)

        response = self.client.post(
            reverse("admin:accounts_user_add"),
            {
                "username": "second-teacher",
                "password1": "safe-test-password",
                "password2": "safe-test-password",
                "role": get_user_model().Role.COURSE_TEACHER,
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("email", response.context["adminform"].form.errors)
        self.assertFalse(get_user_model().objects.filter(username="second-teacher").exists())

    def test_admin_add_link_redirects_to_teacher_creation(self):
        self.client.force_login(self.lead)
        response = self.client.get(reverse("admin:accounts_studentstaffassignment_add"))
        self.assertRedirects(response, reverse("admin:accounts_studentstaffassignment_teacher_add"))















class LearningOutcomeRatingWorkflowTests(TestCase):
    def setUp(self):
        from .models import CoursePlan, IGPPlan, LearningPerformance, SemesterPlan

        self.account = get_user_model().objects.create_user(
            username="rating-teacher", email="rating@example.edu.tw", password="safe-test-password",
            role=get_user_model().Role.COURSE_TEACHER, is_approved=True, is_staff=True,
        )
        self.teacher = Teacher.objects.create(full_name="Rating Teacher", account=self.account)
        self.students = [Student.objects.create(full_name=name) for name in ("Student A", "Student B")]
        self.performances = []
        for student in self.students:
            plan = IGPPlan.objects.create(student=student, academic_year="115", overall_goal="goal")
            semester = SemesterPlan.objects.create(igp_plan=plan, semester=1, course_needs_assessment="needs", goals="semester goal")
            course = CoursePlan.objects.create(semester_plan=semester, course_name="Mathematics", teacher=self.teacher, goals="course goal")
            self.performances.append(LearningPerformance.objects.create(course_plan=course, description=f"Performance {student.pk}"))
        self.source = self.performances[0].course_plan

    def test_course_teacher_rates_all_assigned_students_by_subject(self):
        from .models import LearningOutcomeRating

        self.client.force_login(self.account)
        dashboard = self.client.get(reverse("admin:accounts_learningoutcomerating_changelist"))
        subject_url = reverse("admin:accounts_learningoutcomerating_subject", args=[self.source.pk])
        subject = self.client.get(subject_url)
        response = self.client.post(subject_url, {
            f"rating-{self.performances[0].pk}": "3",
            f"rating-{self.performances[1].pk}": "4",
        })

        self.assertContains(dashboard, "Mathematics")
        self.assertContains(subject, "Student A")
        self.assertContains(subject, "Student B")
        self.assertContains(subject, "<h2>Student A</h2>", html=True)
        self.assertContains(subject, "<h2>Student B</h2>", html=True)
        self.assertRedirects(response, subject_url)
        self.assertEqual(LearningOutcomeRating.objects.get(learning_performance=self.performances[0]).rating, 3)
        self.assertEqual(LearningOutcomeRating.objects.get(learning_performance=self.performances[1]).rating, 4)

    def test_case_manager_without_a_taught_course_cannot_rate_their_case_student(self):
        case_account = get_user_model().objects.create_user(
            username="case-only", email="case-only@example.edu.tw", password="safe-test-password",
            role=get_user_model().Role.CASE_MANAGER, is_approved=True, is_staff=True,
        )
        case_teacher = Teacher.objects.create(full_name="Case Only", account=case_account)
        StudentStaffAssignment.objects.create(
            student=self.students[0], staff=case_teacher, role=StudentStaffAssignment.Role.CASE_MANAGER,
        )

        from .models import LearningOutcomeRating

        LearningOutcomeRating.objects.create(
            learning_performance=self.performances[0],
            rating=LearningOutcomeRating.Rating.GOOD,
            notes="read-only result",
        )
        self.assertFalse(can_manage_learning_outcomes(case_account))
        self.client.force_login(case_account)
        dashboard = self.client.get(reverse("admin:accounts_learningoutcomerating_changelist"))
        response = self.client.get(reverse("admin:accounts_learningoutcomerating_subject", args=[self.source.pk]))
        student_url = reverse("admin:accounts_learningoutcomerating_student", args=[self.students[0].pk])
        student_page = self.client.get(student_url)
        rejected_post = self.client.post(student_url, {"rating": "4"})

        self.assertContains(dashboard, self.students[0].full_name)
        self.assertEqual(response.status_code, 404)
        self.assertEqual(student_page.status_code, 200)
        self.assertContains(student_page, "○3 良好")
        self.assertContains(student_page, "read-only result")
        self.assertNotContains(student_page, "<select")
        self.assertEqual(rejected_post.status_code, 405)

    def test_case_manager_who_teaches_can_rate_only_their_own_courses(self):
        from .models import CoursePlan, IGPPlan, LearningOutcomeRating, LearningPerformance, SemesterPlan

        self.account.role = get_user_model().Role.CASE_MANAGER
        self.account.save(update_fields=["role"])
        for student in self.students:
            StudentStaffAssignment.objects.create(
                student=student, staff=self.teacher, role=StudentStaffAssignment.Role.CASE_MANAGER,
            )
        other_account = get_user_model().objects.create_user(
            username="other-teacher", email="other-teacher@example.edu.tw", password="safe-test-password",
            role=get_user_model().Role.COURSE_TEACHER, is_approved=True, is_staff=True,
        )
        other_teacher = Teacher.objects.create(full_name="Other Teacher", account=other_account)
        other_student = Student.objects.create(full_name="Other Student")
        StudentStaffAssignment.objects.create(
            student=other_student, staff=self.teacher, role=StudentStaffAssignment.Role.CASE_MANAGER,
        )
        other_plan = IGPPlan.objects.create(student=other_student, academic_year="115", overall_goal="goal")
        other_semester = SemesterPlan.objects.create(igp_plan=other_plan, semester=1, goals="semester goal")
        other_course = CoursePlan.objects.create(
            semester_plan=other_semester, course_name="Mathematics", teacher=other_teacher, goals="course goal",
        )
        other_performance = LearningPerformance.objects.create(course_plan=other_course, description="Other performance")

        self.assertTrue(can_manage_learning_outcomes(self.account, self.students[0]))
        self.assertFalse(can_manage_learning_outcomes(self.account, other_student))
        self.client.force_login(self.account)
        own_url = reverse("admin:accounts_learningoutcomerating_subject", args=[self.source.pk])
        own_page = self.client.get(own_url)
        denied_page = self.client.get(reverse("admin:accounts_learningoutcomerating_subject", args=[other_course.pk]))
        response = self.client.post(own_url, {
            f"rating-{self.performances[0].pk}": "4",
            f"rating-{other_performance.pk}": "4",
        })

        self.assertContains(own_page, "Student A")
        self.assertContains(own_page, "Student B")
        self.assertContains(
            own_page,
            reverse("admin:accounts_learningoutcomerating_student", args=[self.students[0].pk]),
        )
        self.assertNotContains(own_page, "Other Student")
        self.assertEqual(denied_page.status_code, 404)
        self.assertRedirects(response, own_url)
        self.assertEqual(LearningOutcomeRating.objects.get(learning_performance=self.performances[0]).rating, 4)
        self.assertFalse(LearningOutcomeRating.objects.filter(learning_performance=other_performance).exists())


class PlanCopyActionTests(TestCase):
    def setUp(self):
        from .models import CoursePlan, IGPPlan, LearningPerformance, SemesterPlan

        self.lead = get_user_model().objects.create_user(
            username="plan-lead", email="plan-lead@example.edu.tw", password="safe-test-password",
            role=get_user_model().Role.SPECIAL_EDUCATION_LEAD, is_approved=True, is_staff=True,
        )
        self.source_student = Student.objects.create(full_name="Source")
        self.target_student = Student.objects.create(full_name="Target")
        self.annual = IGPPlan.objects.create(student=self.source_student, academic_year="115", overall_goal="annual")
        self.semester = SemesterPlan.objects.create(
            igp_plan=self.annual, semester=1, learning_domains="mathematics", special_needs_courses="creativity", goals="semester",
        )
        self.template = CoursePlan.objects.create(
            semester_plan=self.semester, course_name="Mathematics", goals="course", is_template=True, is_active=False,
        )
        self.course = CoursePlan.objects.create(
            semester_plan=self.semester, template=self.template, course_name="Mathematics", goals="course",
        )
        self.template_performance = LearningPerformance.objects.create(
            course_plan=self.template, description="performance", assessment_methods="written",
        )
        self.performance = LearningPerformance.objects.create(course_plan=self.course, description="performance", assessment_methods="written")

    def test_annual_copy_action_and_course_group_list_render(self):
        self.client.force_login(self.lead)
        annual_response = self.client.post(
            reverse("admin:accounts_igpplan_changelist"),
            {"action": "copy_selected_annual_plan", "_selected_action": self.annual.pk},
        )
        course_response = self.client.get(reverse("admin:accounts_courseplan_changelist"))

        self.assertRedirects(annual_response, reverse("admin:accounts_igpplan_copy", args=[self.annual.pk]))
        self.assertEqual(course_response.status_code, 200)
        self.assertEqual(len(course_response.context["course_groups"]), 1)
        self.assertContains(course_response, "Mathematics")
        self.assertNotContains(course_response, self.source_student.full_name)

    def test_course_plan_change_form_renders(self):
        self.client.force_login(self.lead)

        response = self.client.get(reverse("admin:accounts_courseplan_change", args=[self.course.pk]))

        self.assertEqual(response.status_code, 200)

    def test_course_plan_change_form_displays_learning_performances(self):
        self.client.force_login(self.lead)

        response = self.client.get(reverse("admin:accounts_courseplan_change", args=[self.course.pk]))

        self.assertContains(response, "學習表現")
        self.assertContains(response, self.performance.description)
        self.assertContains(response, "learning_performances-TOTAL_FORMS")

    def test_admin_course_plan_needs_are_saved_and_restored(self):
        self.client.force_login(self.lead)
        prefix = "learning_performances"
        response = self.client.post(reverse("admin:accounts_courseplan_change", args=[self.course.pk]), {
            "semester_plan": self.semester.pk, "course_name": self.course.course_name, "teacher": "", "goals": self.course.goals,
            "cognitive_adjustments": ["加深加廣"], "affective_support": ["情意技能"], "skill_training": ["生活技能"],
            f"{prefix}-TOTAL_FORMS": 1, f"{prefix}-INITIAL_FORMS": 1, f"{prefix}-MIN_NUM_FORMS": 0, f"{prefix}-MAX_NUM_FORMS": 1000,
            f"{prefix}-0-id": self.performance.pk, f"{prefix}-0-course_plan": self.course.pk,
            f"{prefix}-0-description": self.performance.description, f"{prefix}-0-adjustment": self.performance.adjustment, f"{prefix}-0-assessment_methods": [],
            "_save": "Save",
        })

        self.assertEqual(response.status_code, 302)
        self.course.refresh_from_db()
        self.assertEqual(self.course.cognitive_adjustments, "加深加廣")
        self.assertEqual(self.course.affective_support, "情意技能")
        self.assertEqual(self.course.skill_training, "生活技能")

    def test_course_plan_needs_are_saved_and_restored(self):
        from .forms import CoursePlanForm

        form = CoursePlanForm({
            "semester_plan": self.semester.pk, "course_name": "Science", "goals": "goal",
            "cognitive_adjustments": ["加深加廣"],
            "affective_support": ["情意技能"], "skill_training": ["生活技能"],
        })

        self.assertTrue(form.is_valid(), form.errors)
        course = form.save()
        self.assertEqual(course.cognitive_adjustments, "加深加廣")
        self.assertEqual(course.affective_support, "情意技能")
        self.assertEqual(course.skill_training, "生活技能")
        restored = CoursePlanForm(instance=course)["cognitive_adjustments"]
        self.assertEqual(restored.value(), ["加深加廣"])
        self.assertIn("checked", str(restored))

    def test_course_group_add_creates_student_plans_from_one_course(self):
        from .models import CoursePlan

        self.client.force_login(self.lead)
        prefix = "learning_performances"
        response = self.client.post(reverse("admin:accounts_courseplan_group_add"), {
            "academic_year": "115",
            "semester": 2,
            "students": [self.source_student.pk, self.target_student.pk],
            "course_name": "Science Inquiry",
            "teacher": "",
            "goals": "shared science goal",
            "activities": "experiment",
            "learning_domains": [],
            "special_needs_courses": [],
            "cognitive_adjustments": [],
            "affective_support": [],
            "skill_training": [],
            f"{prefix}-TOTAL_FORMS": 1,
            f"{prefix}-INITIAL_FORMS": 0,
            f"{prefix}-MIN_NUM_FORMS": 0,
            f"{prefix}-MAX_NUM_FORMS": 1000,
            f"{prefix}-0-id": "",
            f"{prefix}-0-description": "science performance",
            f"{prefix}-0-adjustment": "individual support",
            f"{prefix}-0-assessment_methods": [],
        })

        plans = CoursePlan.objects.filter(course_name="Science Inquiry", is_active=True)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(plans.count(), 2)
        self.assertEqual(
            set(plans.values_list("semester_plan__igp_plan__student_id", flat=True)),
            {self.source_student.pk, self.target_student.pk},
        )
        self.assertEqual(set(plans.values_list("goals", flat=True)), {"shared science goal"})
        self.assertEqual(
            set(plans.values_list("learning_performances__description", flat=True)),
            {"science performance"},
        )

    def test_only_lead_can_delete_annual_and_semester_plans(self):
        from .models import StudentStaffAssignment, Teacher

        case_manager = get_user_model().objects.create_user(
            username="case-delete", email="case-delete@example.edu.tw", password="safe-test-password",
            role=get_user_model().Role.CASE_MANAGER, is_approved=True, is_staff=True,
        )
        teacher = Teacher.objects.create(full_name="Case Manager", account=case_manager)
        StudentStaffAssignment.objects.create(
            student=self.source_student,
            staff=teacher,
            role=StudentStaffAssignment.Role.CASE_MANAGER,
        )

        self.client.force_login(case_manager)
        self.assertEqual(self.client.get(reverse("admin:accounts_igpplan_delete", args=[self.annual.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse("admin:accounts_semesterplan_delete", args=[self.semester.pk])).status_code, 403)

        self.client.force_login(self.lead)
        self.assertEqual(self.client.get(reverse("admin:accounts_igpplan_delete", args=[self.annual.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse("admin:accounts_semesterplan_delete", args=[self.semester.pk])).status_code, 200)

    def test_deleting_template_anchor_reanchors_shared_template(self):
        from .models import CoursePlan, IGPPlan, SemesterPlan

        target_annual = IGPPlan.objects.create(student=self.target_student, academic_year="115", overall_goal="target")
        target_semester = SemesterPlan.objects.create(igp_plan=target_annual, semester=1, goals="target")
        CoursePlan.objects.create(
            semester_plan=target_semester,
            template=self.template,
            course_name="Mathematics",
            goals="target course",
        )
        self.client.force_login(self.lead)

        response = self.client.post(
            reverse("admin:accounts_igpplan_delete", args=[self.annual.pk]),
            {"post": "yes"},
        )

        self.assertEqual(response.status_code, 302)
        self.template.refresh_from_db()
        self.assertEqual(self.template.semester_plan, target_semester)
    def test_adding_student_preserves_individual_version_until_explicit_overwrite(self):
        from .models import CoursePlan, IGPPlan, LearningPerformance, SemesterPlan

        self.client.force_login(self.lead)
        third_student = Student.objects.create(full_name="Third")
        target_annual = IGPPlan.objects.create(student=self.target_student, academic_year="115", overall_goal="target")
        target_semester = SemesterPlan.objects.create(igp_plan=target_annual, semester=1, goals="target semester")
        target = CoursePlan.objects.create(
            semester_plan=target_semester,
            template=self.template,
            course_name="Mathematics",
            goals="individual target goal",
        )
        target_performance = LearningPerformance.objects.create(
            course_plan=target,
            description="individual performance",
        )
        prefix = "learning_performances"
        data = {
            "academic_year": "115",
            "semester": 1,
            "students": [self.source_student.pk, self.target_student.pk, third_student.pk],
            "course_name": "Mathematics",
            "teacher": "",
            "goals": "updated template goal",
            "learning_domains": [],
            "special_needs_courses": [],
            "cognitive_adjustments": [],
            "affective_support": [],
            "skill_training": [],
            f"{prefix}-TOTAL_FORMS": 1,
            f"{prefix}-INITIAL_FORMS": 1,
            f"{prefix}-MIN_NUM_FORMS": 0,
            f"{prefix}-MAX_NUM_FORMS": 1000,
            f"{prefix}-0-id": self.template_performance.pk,
            f"{prefix}-0-course_plan": self.template.pk,
            f"{prefix}-0-description": "updated template performance",
            f"{prefix}-0-adjustment": "",
            f"{prefix}-0-assessment_methods": [],
            "action": "save",
        }

        self.client.post(reverse("admin:accounts_courseplan_group", args=[self.template.pk]), data)
        target.refresh_from_db()
        target_performance.refresh_from_db()
        added = CoursePlan.objects.get(template=self.template, semester_plan__igp_plan__student=third_student)
        self.assertEqual(target.goals, "individual target goal")
        self.assertEqual(target_performance.description, "individual performance")
        self.assertEqual(added.goals, "updated template goal")
        self.assertEqual(added.learning_performances.get().description, "updated template performance")

        data["action"] = "overwrite"
        self.client.post(reverse("admin:accounts_courseplan_group", args=[self.template.pk]), data)
        target.refresh_from_db()
        target_performance.refresh_from_db()
        self.assertEqual(target.goals, "updated template goal")
        self.assertEqual(target_performance.description, "updated template performance")
    def test_learning_performance_uses_automatic_sequence(self):
        from .models import LearningPerformance

        first = LearningPerformance.objects.create(course_plan=self.course, adjustment="first")
        second = LearningPerformance.objects.create(course_plan=self.course, adjustment="second")

        self.assertEqual((first.sort_order, second.sort_order), (2, 3))

    def test_course_group_syncs_students_and_archives_without_deleting_results(self):
        from .models import CoursePlan, LearningOutcomeRating

        self.client.force_login(self.lead)
        group_url = reverse("admin:accounts_courseplan_group", args=[self.course.pk])
        prefix = "learning_performances"
        shared_data = {
            "academic_year": "115",
            "semester": 1,
            "students": [self.source_student.pk, self.target_student.pk],
            "course_name": "Mathematics",
            "teacher": "",
            "goals": "shared course goal",
            "activities": "shared activity",
            "learning_domains": [],
            "special_needs_courses": [],
            "cognitive_adjustments": [],
            "affective_support": [],
            "skill_training": [],
            f"{prefix}-TOTAL_FORMS": 1,
            f"{prefix}-INITIAL_FORMS": 1,
            f"{prefix}-MIN_NUM_FORMS": 0,
            f"{prefix}-MAX_NUM_FORMS": 1000,
            f"{prefix}-0-id": self.template_performance.pk,
            f"{prefix}-0-course_plan": self.template.pk,
            f"{prefix}-0-description": "shared performance",
            f"{prefix}-0-adjustment": "shared adjustment",
            f"{prefix}-0-assessment_methods": [],
        }
        response = self.client.post(group_url, shared_data)

        target = CoursePlan.objects.get(
            semester_plan__igp_plan__student=self.target_student,
            course_name="Mathematics",
        )
        self.assertRedirects(response, reverse("admin:accounts_courseplan_group", args=[self.course.pk]))
        self.assertEqual(target.goals, "shared course goal")
        self.assertEqual(target.semester_plan.learning_domains, "mathematics")
        self.assertEqual(target.learning_performances.get().description, "shared performance")
        group_page = self.client.get(group_url)
        self.assertContains(group_page, self.source_student.full_name)
        self.assertContains(group_page, self.target_student.full_name)
        self.assertContains(group_page, reverse("admin:accounts_courseplan_change", args=[target.pk]))
        self.assertContains(group_page, 'id="add-performance"')
        self.assertContains(group_page, "__prefix__")

        target_performance = target.learning_performances.get()
        rating = LearningOutcomeRating.objects.create(
            learning_performance=target_performance,
            rating=LearningOutcomeRating.Rating.EXCELLENT,
        )
        shared_data["students"] = [self.source_student.pk]
        archive_response = self.client.post(group_url, shared_data)
        target.refresh_from_db()

        self.assertEqual(archive_response.status_code, 302)
        self.assertFalse(target.is_active)
        self.assertTrue(LearningOutcomeRating.objects.filter(pk=rating.pk).exists())


class CounselingRecordWorkflowTests(TestCase):
    def setUp(self):
        self.lead = get_user_model().objects.create_user(
            username="counseling-lead", email="counseling-lead@example.edu.tw",
            password="safe-test-password", role=get_user_model().Role.SPECIAL_EDUCATION_LEAD,
            is_approved=True, is_staff=True,
        )
        self.case_manager = get_user_model().objects.create_user(
            username="counseling-case", email="counseling-case@example.edu.tw",
            password="safe-test-password", role=get_user_model().Role.CASE_MANAGER,
            is_approved=True, is_staff=True,
        )
        self.hidden_case_manager = get_user_model().objects.create_user(
            username="other-case", email="other-case@example.edu.tw",
            password="safe-test-password", role=get_user_model().Role.CASE_MANAGER,
            is_approved=True, is_staff=True,
        )
        self.student = Student.objects.create(full_name="輔導學生")
        self.hidden_student = Student.objects.create(full_name="不可見輔導學生")
        teacher = Teacher.objects.create(full_name="個管教師", account=self.case_manager)
        StudentStaffAssignment.objects.create(
            student=self.student, staff=teacher, role=StudentStaffAssignment.Role.CASE_MANAGER,
        )
        self.list_url = reverse("admin:accounts_counselingrecord_changelist")

    def records_url(self, student):
        return reverse("admin:accounts_counselingrecord_student_records", args=[student.pk])

    def create_record(self):
        self.client.force_login(self.case_manager)
        response = self.client.post(reverse("admin:accounts_counselingrecord_add"), {
            "student": self.student.pk,
            "academic_year": "115",
            "recorded_on": "2026-07-15",
            "event": "學習適應",
            "summary": "與學生討論學習安排。",
            "_save": "儲存",
        })
        self.assertEqual(response.status_code, 302)
        from .models import CounselingRecord
        return CounselingRecord.objects.get()

    def run_action(self, user, action, record):
        self.client.force_login(user)
        return self.client.post(self.records_url(record.student), {
            "action": action,
            "_selected_action": [record.pk],
        })

    def test_case_manager_can_create_and_only_view_assigned_student_records(self):
        record = self.create_record()
        from .models import CounselingRecord
        CounselingRecord.objects.create(
            student=self.hidden_student, author=self.hidden_case_manager,
            event="不可見", summary="不應出現。",
        )

        self.client.force_login(self.case_manager)
        student_list = self.client.get(self.list_url)
        record_list = self.client.get(self.records_url(self.student))
        hidden_list = self.client.get(self.records_url(self.hidden_student))

        self.assertEqual(student_list.status_code, 200)
        self.assertContains(student_list, self.student.full_name)
        self.assertNotContains(student_list, self.hidden_student.full_name)
        self.assertEqual(record_list.status_code, 200)
        for heading in ("日期", "參與人員", "事件", "內容概要敘述", "處遇方式", "記錄人員"):
            self.assertContains(record_list, heading)
        self.assertContains(record_list, record.event)
        self.assertContains(record_list, reverse("admin:accounts_counselingrecord_change", args=[record.pk]))
        self.assertEqual(hidden_list.status_code, 404)
    def test_submit_return_review_and_lock_create_audit_events(self):
        record = self.create_record()
        from .models import AuditEvent, CounselingRecord

        self.assertEqual(record.status, CounselingRecord.Status.DRAFT)
        self.run_action(self.case_manager, "submit_selected", record)
        record.refresh_from_db()
        self.assertEqual(record.status, CounselingRecord.Status.SUBMITTED)

        self.run_action(self.lead, "return_selected", record)
        record.refresh_from_db()
        self.assertEqual(record.status, CounselingRecord.Status.RETURNED)

        self.run_action(self.case_manager, "submit_selected", record)
        self.run_action(self.lead, "review_selected", record)
        record.refresh_from_db()
        self.assertEqual(record.status, CounselingRecord.Status.REVIEWED)
        self.assertEqual(record.reviewed_by, self.lead)

        self.run_action(self.lead, "lock_selected", record)
        record.refresh_from_db()
        self.assertEqual(record.status, CounselingRecord.Status.LOCKED)
        self.assertEqual(record.locked_by, self.lead)
        self.assertEqual(
            list(AuditEvent.objects.filter(target_pk=str(record.pk)).values_list("event_type", flat=True)),
            [
                AuditEvent.EventType.COUNSELING_LOCKED,
                AuditEvent.EventType.COUNSELING_REVIEWED,
                AuditEvent.EventType.COUNSELING_SUBMITTED,
                AuditEvent.EventType.COUNSELING_RETURNED,
                AuditEvent.EventType.COUNSELING_SUBMITTED,
                AuditEvent.EventType.COUNSELING_CREATED,
            ],
        )

    def test_case_manager_cannot_edit_submitted_record_or_review_it(self):
        record = self.create_record()
        self.run_action(self.case_manager, "submit_selected", record)

        self.client.force_login(self.case_manager)
        response = self.client.get(reverse("admin:accounts_counselingrecord_change", args=[record.pk]))
        review_response = self.run_action(self.case_manager, "review_selected", record)

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'name="summary"')
        self.assertEqual(review_response.status_code, 403)

    def test_locked_record_and_audit_events_are_read_only(self):
        record = self.create_record()
        self.run_action(self.case_manager, "submit_selected", record)
        self.run_action(self.lead, "review_selected", record)
        self.run_action(self.lead, "lock_selected", record)
        from .models import AuditEvent

        self.client.force_login(self.lead)
        record_response = self.client.get(reverse("admin:accounts_counselingrecord_change", args=[record.pk]))
        audit_response = self.client.get(reverse("admin:accounts_auditevent_change", args=[AuditEvent.objects.first().pk]))

        self.assertEqual(record_response.status_code, 200)
        self.assertNotContains(record_response, 'name="summary"')
        self.assertEqual(audit_response.status_code, 200)
        self.assertNotContains(audit_response, 'name="summary"')

    def test_homeroom_and_course_teachers_follow_their_counseling_scopes(self):
        from .models import CounselingRecord, CoursePlan, IGPPlan, SemesterPlan
        from .policies import (
            can_add_counseling_record,
            can_edit_counseling_record,
            counseling_records_for,
        )

        homeroom = get_user_model().objects.create_user(
            username="counseling-homeroom", email="counseling-homeroom@example.edu.tw",
            password="safe-test-password", role=get_user_model().Role.HOMEROOM_TEACHER,
            is_approved=True, is_staff=True,
        )
        course_teacher = get_user_model().objects.create_user(
            username="counseling-course", email="counseling-course@example.edu.tw",
            password="safe-test-password", role=get_user_model().Role.COURSE_TEACHER,
            is_approved=True, is_staff=True,
        )
        homeroom_teacher = Teacher.objects.create(full_name="班級導師", account=homeroom)
        course_teacher_record = Teacher.objects.create(full_name="任課教師", account=course_teacher)
        StudentStaffAssignment.objects.create(
            student=self.student, staff=homeroom_teacher,
            role=StudentStaffAssignment.Role.HOMEROOM_TEACHER,
        )
        annual = IGPPlan.objects.create(student=self.student, academic_year="115", overall_goal="goal")
        semester = SemesterPlan.objects.create(igp_plan=annual, semester=1, goals="goal")
        CoursePlan.objects.create(
            semester_plan=semester, course_name="數學", teacher=course_teacher_record, goals="goal",
        )
        case_record = CounselingRecord.objects.create(
            student=self.student, author=self.case_manager, event="個管紀錄", summary="summary",
        )
        homeroom_record = CounselingRecord.objects.create(
            student=self.student, author=homeroom, event="導師紀錄", summary="summary",
        )
        course_record = CounselingRecord.objects.create(
            student=self.student, author=course_teacher, event="任課紀錄", summary="summary",
        )

        self.assertSetEqual(
            set(counseling_records_for(homeroom).values_list("pk", flat=True)),
            {case_record.pk, homeroom_record.pk, course_record.pk},
        )
        self.assertTrue(can_add_counseling_record(homeroom, self.student))
        self.assertFalse(can_add_counseling_record(homeroom, self.hidden_student))
        self.assertTrue(can_edit_counseling_record(homeroom, homeroom_record))
        self.assertFalse(can_edit_counseling_record(homeroom, case_record))

        self.client.force_login(homeroom)
        homeroom_list = self.client.get(self.records_url(self.student))
        own_change = self.client.get(reverse("admin:accounts_counselingrecord_change", args=[homeroom_record.pk]))
        other_change = self.client.get(reverse("admin:accounts_counselingrecord_change", args=[case_record.pk]))
        self.assertContains(homeroom_list, case_record.event)
        self.assertContains(homeroom_list, homeroom_record.event)
        self.assertEqual(own_change.status_code, 200)
        self.assertContains(own_change, 'name="summary"')
        self.assertEqual(other_change.status_code, 200)
        self.assertNotContains(other_change, 'name="summary"')

        self.assertSetEqual(
            set(counseling_records_for(course_teacher).values_list("pk", flat=True)),
            {course_record.pk},
        )
        self.assertTrue(can_add_counseling_record(course_teacher, self.student))
        self.assertFalse(can_add_counseling_record(course_teacher, self.hidden_student))
        self.assertTrue(can_edit_counseling_record(course_teacher, course_record))
        self.assertFalse(can_edit_counseling_record(course_teacher, case_record))

        self.client.force_login(course_teacher)
        add_page = self.client.get(reverse("admin:accounts_counselingrecord_add"))
        self.assertContains(add_page, self.student.full_name)
        self.assertNotContains(add_page, self.hidden_student.full_name)

    def test_case_manager_can_author_course_students_without_viewing_others_records(self):
        from django.contrib import admin as django_admin
        from django.test import RequestFactory

        from .admin import CoursePlanAdmin
        from .models import CounselingRecord, CoursePlan, IGPPlan, SemesterPlan
        from .policies import counseling_records_for, students_for_course_plans

        case_teacher = Teacher.objects.get(account=self.case_manager)
        course_students = [self.student, self.hidden_student]
        course_students.extend(Student.objects.create(full_name=f"任課學生 {index}") for index in range(3, 5))
        for student in course_students:
            annual = IGPPlan.objects.create(student=student, academic_year="115", overall_goal="goal")
            semester = SemesterPlan.objects.create(igp_plan=annual, semester=1, goals="goal")
            CoursePlan.objects.create(
                semester_plan=semester, course_name="跨領域專題", teacher=case_teacher, goals="goal",
            )

        own_course_record = CounselingRecord.objects.create(
            student=self.hidden_student, author=self.case_manager,
            event="任課學生紀錄", summary="個管教師兼任授課教師的紀錄。",
        )
        other_record = CounselingRecord.objects.create(
            student=self.hidden_student, author=self.hidden_case_manager,
            event="他人紀錄", summary="不應被個管教師甲讀取。",
        )

        self.assertSetEqual(
            set(students_for_course_plans(self.case_manager).values_list("pk", flat=True)),
            {student.pk for student in course_students},
        )
        self.assertIn(own_course_record.pk, counseling_records_for(self.case_manager).values_list("pk", flat=True))
        self.assertNotIn(other_record.pk, counseling_records_for(self.case_manager).values_list("pk", flat=True))

        request = RequestFactory().get("/admin/accounts/courseplan/")
        request.user = self.case_manager
        course_admin = CoursePlanAdmin(CoursePlan, django_admin.site)
        self.assertSetEqual(
            set(course_admin.get_queryset(request).values_list("semester_plan__igp_plan__student_id", flat=True)),
            {student.pk for student in course_students},
        )

        off_course_student = Student.objects.create(full_name="未任課學生")
        self.client.force_login(self.case_manager)
        add_page = self.client.get(reverse("admin:accounts_counselingrecord_add"))
        self.assertContains(add_page, self.hidden_student.full_name)
        self.assertNotContains(add_page, off_course_student.full_name)
        change_page = self.client.get(reverse("admin:accounts_counselingrecord_change", args=[own_course_record.pk]))
        self.assertEqual(change_page.status_code, 200)
        self.assertContains(change_page, 'name="summary"')

    def test_actual_case_assignment_is_honored_for_a_different_primary_role(self):
        from .models import CounselingRecord
        from .policies import (
            counseling_records_for,
            students_for_counseling_authoring,
            students_for_counseling_index,
        )

        multi_role_user = get_user_model().objects.create_user(
            username="multi-role-case", email="multi-role-case@example.edu.tw",
            password="safe-test-password", role=get_user_model().Role.COURSE_TEACHER,
            is_approved=True, is_staff=True,
        )
        multi_role_teacher = Teacher.objects.create(full_name="兼任個管教師", account=multi_role_user)
        StudentStaffAssignment.objects.create(
            student=self.hidden_student,
            staff=multi_role_teacher,
            role=StudentStaffAssignment.Role.CASE_MANAGER,
        )
        other_record = CounselingRecord.objects.create(
            student=self.hidden_student,
            author=self.hidden_case_manager,
            event="既有個管紀錄",
            summary="多重身分個管應可閱讀。",
        )

        self.assertIn(self.hidden_student, students_for_counseling_authoring(multi_role_user))
        self.assertIn(self.hidden_student, students_for_counseling_index(multi_role_user))
        self.assertIn(other_record, counseling_records_for(multi_role_user))

        self.client.force_login(multi_role_user)
        student_list = self.client.get(self.list_url)
        record_list = self.client.get(self.records_url(self.hidden_student))
        add_page = self.client.get(reverse("admin:accounts_counselingrecord_add") + f"?student={self.hidden_student.pk}")

        self.assertContains(student_list, self.hidden_student.full_name)
        self.assertContains(record_list, other_record.event)
        self.assertContains(add_page, f'value="{self.hidden_student.pk}" selected')
    def test_counseling_form_saves_required_columns_and_lead_can_optionally_add_review_note(self):
        self.client.force_login(self.case_manager)
        response = self.client.post(reverse("admin:accounts_counselingrecord_add"), {
            "student": self.student.pk,
            "academic_year": "115",
            "recorded_on": "2026-07-15",
            "participants": ["本人", "家長", "原班導師", "個管老師", "資優任課"],
            "participants_other": "輔導教師",
            "event": "親師晤談",
            "summary": "討論近期學習壓力與作業安排。",
            "intervention": ["定期晤談", "持續觀察"],
            "intervention_other": "安排時間管理練習",
            "_save": "儲存",
        })
        self.assertEqual(response.status_code, 302)
        from .models import CounselingRecord
        record = CounselingRecord.objects.get()
        self.assertEqual(record.participants, "本人、家長、原班導師、個管老師、資優任課、其他：輔導教師")
        self.assertEqual(record.event, "親師晤談")
        self.assertEqual(record.summary, "討論近期學習壓力與作業安排。")
        self.assertEqual(record.intervention, "定期晤談、持續觀察、其他：安排時間管理練習")

        self.run_action(self.case_manager, "submit_selected", record)
        self.client.force_login(self.lead)
        change_url = reverse("admin:accounts_counselingrecord_change", args=[record.pk])
        change_page = self.client.get(change_url)
        self.assertContains(change_page, 'name="review_note"')
        response = self.client.post(change_url, {"review_note": "建議持續追蹤。", "_save": "儲存"})
        self.assertEqual(response.status_code, 302)
        record.refresh_from_db()
        self.assertEqual(record.review_note, "建議持續追蹤。")
